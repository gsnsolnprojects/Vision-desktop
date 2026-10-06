from visionm_engine.exporter import build_run_report
from visionm_engine.run_store import RunStore

MODEL = {"id": "m1", "name": "Model 1", "modelType": "YOLO", "backendModelId": "m1", "company": "c", "project": "p"}


def make_store(tmp_path):
    return RunStore(tmp_path)


def test_run_lifecycle_and_sync_bookkeeping(tmp_path):
    s = make_store(tmp_path)
    run_id = s.create_run(model=MODEL, user={"id": "u"}, camera_name="cam", config={"live": {"goodClasses": ["product"]}}, sync=True)
    d = s.add_defect(run_id, ["dent"], b"\xff\xd8fakejpeg")
    s.finish_run(run_id, {"totalProducts": 2, "goodProducts": 1, "badProducts": 1, "classCounts": {"product": 2, "dent": 1}}, {})

    run = s.get_run(run_id)
    assert run["status"] == "completed" and run["defectCount"] == 1
    assert [r["id"] for r in s.runs_needing_sync()] == [run_id]

    s.mark_run_synced(run_id, "backend-1", final=True)
    # still has an unsynced defect -> still needs sync
    assert [r["id"] for r in s.runs_needing_sync()] == [run_id]
    s.mark_defect_synced(d["id"])
    assert s.runs_needing_sync() == []


def test_local_only_runs_never_sync(tmp_path):
    s = make_store(tmp_path)
    run_id = s.create_run(model=MODEL, user=None, camera_name="cam", config={}, sync=False)
    s.finish_run(run_id, {}, {})
    assert s.runs_needing_sync() == []


def test_crash_recovery_marks_running_runs_interrupted(tmp_path):
    s = make_store(tmp_path)
    run_id = s.create_run(model=MODEL, user=None, camera_name="cam", config={}, sync=False)
    s.close()
    s2 = make_store(tmp_path)
    assert s2.get_run(run_id)["status"] == "interrupted"


def test_excel_report(tmp_path):
    import cv2
    import numpy as np

    s = make_store(tmp_path)
    run_id = s.create_run(model=MODEL, user={"email": "op@x"}, camera_name="cam", config={"live": {"goodClasses": ["product"]}}, sync=False)
    ok, jpg = cv2.imencode(".jpg", np.zeros((720, 1280, 3), np.uint8))
    s.add_defect(run_id, ["dent"], jpg.tobytes())
    s.finish_run(run_id, {"totalProducts": 1, "goodProducts": 0, "badProducts": 1, "classCounts": {"product": 1, "dent": 1}}, {})
    data = build_run_report(s, run_id)
    assert data[:2] == b"PK"  # xlsx is a zip

    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Summary", "LiveDefectLogs"]
    assert wb["LiveDefectLogs"]["D2"].value == "dent"
    assert len(wb["LiveDefectLogs"]._images) == 1


def test_run_name_and_notes(tmp_path):
    s = make_store(tmp_path)
    run_id = s.create_run(model=MODEL, user=None, camera_name="cam", config={}, sync=False)
    assert s.get_run(run_id)["name"] is None
    assert s.update_meta(run_id, "  Morning shift  ", "Line 1, batch 42")
    r = s.get_run(run_id)
    assert (r["name"], r["notes"]) == ("Morning shift", "Line 1, batch 42")
    assert s.update_meta(run_id, "", None) and s.get_run(run_id)["name"] is None
    assert not s.update_meta("missing", "x", None)


def test_migration_adds_columns_to_old_db(tmp_path):
    import sqlite3

    db = sqlite3.connect(tmp_path / "station.db")
    db.executescript("CREATE TABLE runs (id TEXT PRIMARY KEY, started_at TEXT NOT NULL, ended_at TEXT, status TEXT NOT NULL, model_id TEXT NOT NULL, model_name TEXT NOT NULL, model_type TEXT, backend_model_id TEXT, company TEXT, project TEXT, project_id TEXT, user_json TEXT, camera_name TEXT, config_json TEXT NOT NULL, analytics_json TEXT NOT NULL, stats_json TEXT NOT NULL DEFAULT '{}', sync_state TEXT NOT NULL DEFAULT 'pending', sync_error TEXT, backend_run_id TEXT);")
    db.close()
    s = make_store(tmp_path)
    run_id = s.create_run(model=MODEL, user=None, camera_name="cam", config={}, sync=False)
    assert s.update_meta(run_id, "ok", None)
