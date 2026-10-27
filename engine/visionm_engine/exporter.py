"""Excel report for a run: summary + defect log with snapshots (same columns as the web app)."""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

import cv2

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill


from .annotate import make_thumbnail
from .run_store import RunStore

THUMB_W, THUMB_H = 160, 90
_HEADER_FILL = PatternFill("solid", fgColor="E5E7EB")


def _local(ts: str | None) -> str:
    if not ts:
        return ""
    try:
        return datetime.fromisoformat(ts).astimezone().strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return ts


def _xl_image(data: bytes) -> XlImage:
    xl = XlImage(io.BytesIO(data))
    ratio = min(THUMB_W / xl.width, THUMB_H / xl.height)
    xl.width, xl.height = int(xl.width * ratio), int(xl.height * ratio)
    return xl


def _thumb_from_file(path: Path) -> bytes | None:
    """Older snapshots without a stored thumbnail: decode at 1/4 size (fast) and shrink."""
    img = cv2.imread(str(path), cv2.IMREAD_REDUCED_COLOR_4)
    return make_thumbnail(img) if img is not None else None


def build_run_report(store: RunStore, run_id: str) -> bytes:
    run = store.get_run(run_id)
    if run is None:
        raise KeyError(run_id)
    a = run["analytics"] or {}
    good_classes = set((run["config"].get("live") or {}).get("goodClasses") or [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    total = a.get("totalProducts", 0)
    bad = a.get("badProducts", 0)
    rows = [
        ("Run", run.get("name") or run["id"]),
        ("Run ID", run["id"]),
        ("Notes", run.get("notes") or "-"),
        ("Model", run["model"]["name"]),
        ("Project", run.get("project") or "-"),
        ("Camera", run.get("cameraName") or "-"),
        ("Operator", (run.get("user") or {}).get("email") or "-"),
        ("Started", _local(run["startedAt"])),
        ("Ended", _local(run["endedAt"])),
        ("Status", run["status"]),
        ("", ""),
        ("Total products", total),
        ("Good products", a.get("goodProducts", 0)),
        ("Bad products", bad),
        ("Defect rate", f"{(bad / total * 100):.1f}%" if total else "0.0%"),
        ("Defect log entries", run["defectCount"]),
        ("Avg inference time (ms)", (run.get("stats") or {}).get("avgInferMs") or "-"),
    ]
    for r in rows:
        ws.append(list(r))
    for c in ws["A"]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 44

    ws.append([])
    ws.append(["Defect class", "Products affected", "Share of bad products"])
    hdr = ws.max_row
    for c in ws[hdr]:
        c.font, c.fill = Font(bold=True), _HEADER_FILL
    defect_counts = sorted(
        ((k, v) for k, v in (a.get("classCounts") or {}).items() if k not in good_classes),
        key=lambda kv: kv[1],
        reverse=True,
    )
    for cls, n in defect_counts:
        ws.append([cls, n, f"{(n / bad * 100):.1f}%" if bad else "0.0%"])
    ws.column_dimensions["C"].width = 22

    logs = wb.create_sheet("LiveDefectLogs")
    headers = ["SrNo", "Timestamp", "Verdict", "DefectClasses", "LoggedFrom", "Snapshot"]
    widths = [8, 24, 12, 28, 18, 26]
    logs.append(headers)
    for i, w in enumerate(widths):
        logs.column_dimensions[chr(65 + i)].width = w
    for c in logs[1]:
        c.font, c.fill = Font(bold=True), _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    logs.row_dimensions[1].height = 24

    defects = list(reversed(store.list_defects(run_id)))  # oldest first, like a log
    thumbs = store.run_thumbs(run_id)
    middle, wrap = Alignment(vertical="center"), Alignment(vertical="center", wrap_text=True)
    for idx, d in enumerate(defects, start=1):
        row = idx + 1
        values = [idx, _local(d["timestamp"]), d["verdict"], ", ".join(d["defectClasses"]), "Live Camera", ""]
        for col, v in enumerate(values, start=1):  # direct cell access: ws[row] rescans the whole sheet each time
            c = logs.cell(row=row, column=col, value=v)
            c.alignment = wrap if col == 4 else middle
        logs.row_dimensions[row].height = 72
        data = thumbs.get(d["id"])
        if data is None and d["hasImage"]:
            path = store.defect_image_path(run_id, d["id"])
            data = _thumb_from_file(path) if path else None
        if data is not None:
            logs.add_image(_xl_image(data), f"F{row}")
        elif not d["hasImage"]:
            logs.cell(row=row, column=6, value="Snapshot removed")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
